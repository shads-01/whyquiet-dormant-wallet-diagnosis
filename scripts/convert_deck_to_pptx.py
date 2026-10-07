# pyright: reportMissingImports=false
import os

from bs4 import BeautifulSoup
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


def build_pptx(html_path, output_pptx):
    with open(html_path, "r", encoding="utf-8") as f:
        soup = BeautifulSoup(f.read(), "html.parser")

    prs = Presentation()
    # 16:9 Widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6]

    # Colors
    BG_COLOR = RGBColor(250, 250, 250)      # #FAFAFA
    CARD_BG = RGBColor(255, 255, 255)       # #FFFFFF
    TEXT_MAIN = RGBColor(17, 20, 24)        # #111418
    TEXT_MUTED = RGBColor(100, 116, 139)    # #64748B
    BORDER_COLOR = RGBColor(226, 232, 240)  # #E2E8F0
    
    ACCENT_RED = RGBColor(224, 83, 56)      # #E05338
    ACCENT_BLUE = RGBColor(37, 99, 235)     # #2563EB
    ACCENT_GREEN = RGBColor(16, 185, 129)   # #10B981

    slides_data = soup.select(".slide")
    
    for idx, s in enumerate(slides_data, start=1):
        slide = prs.slides.add_slide(blank_slide_layout)
        
        # Background fill
        bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = BG_COLOR
        bg_shape.line.fill.background()

        # Parse Elements
        tag_elem = s.select_one(".category-tag")
        headline_elem = s.select_one(".slide-headline")
        lead_elem = s.select_one(".lead-paragraph")
        img_elem = s.select_one("img")
        stat_cards = s.select(".stat-card")
        pill_items = s.select(".pill-item")

        # Determine layout direction (Reverse = image on left)
        is_reverse = "reverse" in s.select_one(".slide-body").get("class", []) if s.select_one(".slide-body") else False

        # 1. Header (Category Tag + Headline)
        header_box = slide.shapes.add_textbox(Inches(0.66), Inches(0.5), Inches(12.0), Inches(1.3))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        # Category Tag
        if tag_elem:
            p_tag = tf.paragraphs[0]
            p_tag.text = tag_elem.text.strip().upper()
            p_tag.font.size = Pt(11)
            p_tag.font.bold = True
            p_tag.font.name = "Arial"
            if "blue" in tag_elem.get("class", []):
                p_tag.font.color.rgb = ACCENT_BLUE
            elif "green" in tag_elem.get("class", []):
                p_tag.font.color.rgb = ACCENT_GREEN
            else:
                p_tag.font.color.rgb = ACCENT_RED
            p_tag.space_after = Pt(4)

        # Headline
        if headline_elem:
            p_head = tf.add_paragraph() if tag_elem else tf.paragraphs[0]
            p_head.text = headline_elem.text.strip()
            p_head.font.size = Pt(26)
            p_head.font.bold = True
            p_head.font.color.rgb = TEXT_MAIN
            p_head.font.name = "Arial"

        # 2. Left / Right Column Positions
        if is_reverse:
            # Image on Left, Content on Right
            img_left, img_top, img_width, img_height = Inches(0.66), Inches(2.0), Inches(5.4), Inches(4.5)
            content_left, content_top, content_width = Inches(6.4), Inches(2.0), Inches(6.2)
        else:
            # Content on Left, Image on Right
            content_left, content_top, content_width = Inches(0.66), Inches(2.0), Inches(6.2)
            img_left, img_top, img_width, img_height = Inches(7.2), Inches(2.0), Inches(5.4), Inches(4.5)

        # 3. Add Image
        if img_elem:
            img_src = img_elem.get("src", "")
            img_path = os.path.join(os.path.dirname(html_path), img_src)
            if os.path.exists(img_path):
                # Draw border card behind image
                border_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, img_left, img_top, img_width, img_height)
                border_card.fill.solid()
                border_card.fill.fore_color.rgb = CARD_BG
                border_card.line.color.rgb = TEXT_MAIN
                border_card.line.width = Pt(2)

                # Insert actual image slightly inset
                slide.shapes.add_picture(img_path, img_left + Inches(0.08), img_top + Inches(0.08), width=img_width - Inches(0.16))

        # 4. Content Block (Lead + Stats / Pills)
        content_box = slide.shapes.add_textbox(content_left, content_top, content_width, Inches(4.5))
        ctf = content_box.text_frame
        ctf.word_wrap = True
        ctf.margin_left = ctf.margin_top = ctf.margin_right = ctf.margin_bottom = 0

        # Lead Paragraph
        if lead_elem:
            p_lead = ctf.paragraphs[0]
            p_lead.text = lead_elem.text.strip()
            p_lead.font.size = Pt(14)
            p_lead.font.color.rgb = RGBColor(51, 65, 85)
            p_lead.font.name = "Arial"
            p_lead.space_after = Pt(14)

        # Pill Items
        if pill_items:
            for pill in pill_items:
                p_pill = ctf.add_paragraph()
                p_pill.text = "• " + pill.text.strip()
                p_pill.font.size = Pt(12)
                p_pill.font.bold = True
                p_pill.font.color.rgb = TEXT_MAIN
                p_pill.font.name = "Arial"
                p_pill.space_after = Pt(8)

        # Stat Cards (2x2 grid if present)
        if stat_cards:
            card_top = content_top + Inches(1.3 if lead_elem else 0)
            for c_idx, sc in enumerate(stat_cards):
                num_elem = sc.select_one(".stat-number")
                lbl_elem = sc.select_one(".stat-label")
                sub_elem = sc.select_one(".stat-subtext")

                col_num = c_idx % 2
                row_num = c_idx // 2

                sc_left = content_left + Inches(col_num * 3.0)
                sc_top = card_top + Inches(row_num * 1.5)
                sc_w, sc_h = Inches(2.8), Inches(1.3)

                # Card Background Box
                sc_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, sc_left, sc_top, sc_w, sc_h)
                sc_box.fill.solid()
                sc_box.fill.fore_color.rgb = CARD_BG
                sc_box.line.color.rgb = BORDER_COLOR
                sc_box.line.width = Pt(1)

                # Card Text
                sc_tf = sc_box.text_frame
                sc_tf.word_wrap = True
                sc_tf.margin_left = sc_tf.margin_top = Inches(0.12)

                if num_elem:
                    p_num = sc_tf.paragraphs[0]
                    p_num.text = num_elem.text.strip()
                    p_num.font.size = Pt(22)
                    p_num.font.bold = True
                    p_num.font.name = "Arial"
                    if "accent" in num_elem.get("class", []):
                        p_num.font.color.rgb = ACCENT_RED
                    elif "blue" in num_elem.get("class", []):
                        p_num.font.color.rgb = ACCENT_BLUE
                    elif "green" in num_elem.get("class", []):
                        p_num.font.color.rgb = ACCENT_GREEN
                    else:
                        p_num.font.color.rgb = TEXT_MAIN

                if lbl_elem:
                    p_lbl = sc_tf.add_paragraph()
                    p_lbl.text = lbl_elem.text.strip().upper()
                    p_lbl.font.size = Pt(9)
                    p_lbl.font.bold = True
                    p_lbl.font.color.rgb = TEXT_MUTED

                if sub_elem:
                    p_sub = sc_tf.add_paragraph()
                    p_sub.text = sub_elem.text.strip()
                    p_sub.font.size = Pt(8)
                    p_sub.font.color.rgb = RGBColor(100, 116, 139)

        # 5. Footer (WhyQuiet Pitch Deck | Slide XX / 07)
        footer_box = slide.shapes.add_textbox(Inches(0.66), Inches(6.8), Inches(12.0), Inches(0.4))
        ftf = footer_box.text_frame
        p_ft = ftf.paragraphs[0]
        p_ft.text = f"WhyQuiet Pitch Deck  |  Slide {idx:02d} / {len(slides_data):02d}"
        p_ft.font.size = Pt(10)
        p_ft.font.color.rgb = TEXT_MUTED
        p_ft.font.name = "Arial"

    try:
        prs.save(output_pptx)
        print(f"Successfully exported PPTX to: {output_pptx}")
    except PermissionError:
        fallback = output_pptx.replace(".pptx", "_copy.pptx")
        prs.save(fallback)
        print(f"File locked by PowerPoint. Saved fallback copy to: {fallback}")

if __name__ == "__main__":
    html_file = "docs/slides/deck.html"
    pptx_file = "docs/slides/deck.pptx"
    build_pptx(html_file, pptx_file)
